"""Turn TallyPrime exports into the CSVs the group dashboard reads.

Usage:
    python3 scripts/build_data.py <file> [<file> ...] --company HD-ATL --name "HD Attingal" [--out data]
    python3 scripts/build_data.py --merge-only [--out data]
    python3 scripts/build_data.py <file> ... --profile canon --company CN-TVM --name "Canon Thiruvananthapuram" --out data/canon

--profile picks the business the registers come from: honda (the default: two-wheeler service bill registers) or
canon (Canon office-automation sales registers: machines, printers and scanners, consumables, spares, service
and TG contracts). Each profile keeps its own output folder.

Pass every file one company sent in a single run: its Service Bill Registers (one per financial year)
and its Stock Summaries. Each file is recognised by its contents, so names do not matter. The run
rewrites <out>/companies/<CODE>/ and then the group files in <out>/ (every row carries a `company`
column), plus:
  companies.csv     one row per company: name, register period, bills, gross
  completeness.csv  the six files the BPro data request asks of every company, received or pending

Every register must add up to its own Grand Total and every stock summary to its own, or the run
stops. Registers name and order their ledger columns differently, so columns are matched by name.
Files received that the script does not read (stock movement reports) are recorded by hand in
<out>/received.csv (company,file) - see README.

Tally's ".xls" export is really an .xlsx file, so it is opened by content, not extension.
Only aggregates are written; customer names, phones and vehicle numbers never leave the register.
"""
import argparse
import collections
import csv
import datetime
import glob
import json
import os
import re

SERVICE = {'Service Bill', 'AMC Service Bill'}
TRANSFER = {'Branch Transfer Out Spare', 'Branch Transfer Out Warranty'}
STREAMS = ['Spare parts', 'Accessories', 'Labour & service charges', 'AMC & warranty', 'Branch transfers', 'Other sales']
# six streams, like the Honda ones, so the dashboard's validated six-colour palette covers them; the ledger view keeps
# every head apart
CANON_STREAMS = ['Copiers & machines', 'Printers & scanners', 'Consumables & spares', 'Service & TG contracts',
                 'Branch transfers', 'Other sales']
CANON_SERVICE = {'Service Bill Invoice', 'TG BILL INVOICE'}   # service visits, and TG meter-reading bills
META = {'DATE', 'PARTICULARS', 'VOUCHER TYPE', 'VOUCHER NO.', 'VOUCHER REF. NO.', 'GSTIN/UIN', 'QUANTITY', 'VALUE', 'GROSS TOTAL'}

# The data request playbook (sections 02-03): six files per company.
REQUIRED = [
    ('Sales register', 'FY 2026-27 Apr-Sep', 'SalesRegister_FY2026-27_Apr-Sep'),
    ('Sales register', 'FY 2025-26', 'SalesRegister_FY2025-26'),
    ('Stock summary', 'As at 31 Mar 2026', 'StockSummary_2026-03-31'),
    ('Stock summary', 'As at 30 Sep 2026', 'StockSummary_2026-09-30'),
    ('Stock movement', 'FY 2025-26', 'StockMovement_FY2025-26'),
    ('Stock movement', 'FY 2026-27 Apr-Sep', 'StockMovement_FY2026-27_Apr-Sep'),
]

FILES = {
    'monthly': ['fy', 'month', 'invoices', 'service_jobs', 'gross', 'taxable', 'gst', 'discount', 'service_gross'],
    'streams': ['fy', 'month', 'stream', 'amount'],
    'heads': ['fy', 'head', 'group', 'amount'],
    'voucher_types': ['fy', 'voucher_type', 'invoices', 'gross', 'taxable'],
    'daily': ['fy', 'date', 'invoices', 'service_jobs', 'gross'],
    'items': ['fy', 'part_no', 'item', 'qty', 'value', 'lines', 'main_channel'],
    'stock': ['as_at', 'category', 'items', 'qty', 'value', 'idle_items', 'idle_value'],
    'stock_items': ['as_at', 'part_no', 'item', 'category', 'qty', 'value', 'sold_qty', 'sold_months', 'cover_months'],
}


def num(v):
    return v if isinstance(v, (int, float)) else 0


def clean(h):
    return re.sub(r'\s+', ' ', str(h or '').replace('_x000D_', ' ').replace('\xa0', ' ')).strip()


def fy_of(d):
    """Indian financial year label for a date: 1 Apr 2026 - 31 Mar 2027 is '2026-27'."""
    y = d.year if d.month >= 4 else d.year - 1
    return f'{y}-{(y + 1) % 100:02d}'


def base_type(t):
    """'Service Bill - New' and 'Service Bill' are the same kind of voucher."""
    return re.sub(r'\s*-\s*New$', '', t or '')


def classify(h):
    """Which part of the bill a register column holds. Registers list their ledgers in different orders
    and spellings, so columns are found by name, never by position."""
    u = clean(h).upper()
    if u in META:
        return 'meta'
    if u in ('CGST', 'SGST', 'IGST'):
        return 'gst'
    if u == 'ROUND OFF':
        return 'round'
    if 'DISCOUNT' in u:
        return 'discount'
    if 'TRANSFER' in u:
        return 'Branch transfers'
    if 'SALES' in u.replace('SALE ', 'SALES ') and 'SPARE' in u:
        return 'Spare parts'
    if 'ACCESSOR' in u and 'FITMENT' not in u:
        return 'Accessories'
    if u in ('AMC SALE',) or 'EXTENDED WARRANTY' in u or u.startswith('EW SALE'):
        return 'AMC & warranty'
    if 'SCRAP' in u or 'POLIMER' in u or 'POLYMER' in u or 'MISCELLANEOUS' in u:
        return 'Other sales'
    return 'Labour & service charges'


def classify_canon(h):
    """A Canon register column. TG (cost-per-copy contract) ledgers are matched before the product words they also
    carry ("Sales CIPL TG Consumables", "Service Charge - TG")."""
    u = clean(h).upper()
    if u in META:
        return 'meta'
    if u in ('CGST', 'SGST', 'IGST'):
        return 'gst'
    if u == 'ROUND OFF':
        return 'round'
    if 'DISCOUNT' in u:
        return 'discount'
    if 'TRANSFER' in u:
        return 'Branch transfers'
    if re.search(r'\bTG\b', u):
        return 'Service & TG contracts'
    if 'MACHINE' in u or 'PLOTTER' in u:
        return 'Copiers & machines'
    if 'PRINTER' in u or 'SCANNER' in u:
        return 'Printers & scanners'
    if 'CONSUMABLE' in u or 'CARTRIDGE' in u or 'SPARE' in u:
        return 'Consumables & spares'
    if 'SERVICE CHARGE' in u:
        return 'Service & TG contracts'
    return 'Other sales'   # buyback, postage and courier, and sales ledgers kept under a person's name


# Canon stock sub-groups (under the "CN" group) carry a code and the GST rate: "CM 18 %", "CO18%"
CANON_CATS = {'AC': 'Canon accessories', 'CA': 'Canon cameras', 'CM': 'Canon copier machines', 'CO': 'Canon consumables (toner, drums)',
              'CT': 'Canon cartridges', 'DR': 'Canon document scanners', 'IC': 'Canon printers', 'PR': 'Canon printers',
              'PT': 'Canon plotters', 'SP': 'Canon spares', 'SUPPLYS': 'Canon supplies'}
CANON_GROUPS = {'BENQ': 'BenQ flat panels', 'CANON DEMO STOCK': 'Demo stock', 'CANON LEASE STOCK': 'Lease stock', 'PITNEY BOWES': 'Pitney Bowes'}


def head_name(h):
    """A ledger as published: Canon books some sales to ledgers named after a person ("NAME PLACE(SALES)"); the
    published data shows those as agent sales, without the name."""
    return 'Agent sales' if PROFILE == 'canon' and re.search(r'\(\s*SALES\s*\)', h, re.I) else h


def is_transfer(t):
    return t in TRANSFER or 'TRFR' in (t or '').upper() or 'TRANSFER' in (t or '').upper()


def use_profile(name):
    """Point the summaries at one business's ledgers and vouchers."""
    global STREAMS, SERVICE, classify, PROFILE
    PROFILE = name
    if name == 'canon':
        STREAMS, SERVICE, classify = CANON_STREAMS, CANON_SERVICE, classify_canon


PROFILE = 'honda'


def category(group, sub):
    """The name a stock line is reported under: its group, or for Canon the readable name of its sub-group."""
    if PROFILE != 'canon':
        return group
    g = re.sub(r'\s*\d+\s*%$', '', clean(group)).upper()
    if sub and g == 'CN':
        code = re.sub(r'[\d%\s]+$', '', clean(sub)).split(' ')[0].upper()
        return CANON_CATS.get(code, 'Canon ' + clean(sub).title())
    return CANON_GROUPS.get(g, clean(group).title())


def sheet_rows(path):
    import openpyxl
    ws = openpyxl.load_workbook(open(path, 'rb'), read_only=True, data_only=True).worksheets[0]
    return ws.title, list(ws.iter_rows(values_only=True))


def split_item(text):
    mm = re.match(r'^(.*)\[([^\]]+)\]\s*$', text)
    return (mm.group(1).strip(), mm.group(2).strip()) if mm else (text.strip(), '')


def summarise(path, rows):
    hdr_at = next(i for i, r in enumerate(rows) if r and r[0] == 'Date' and r[1] == 'Particulars')
    hdr = [clean(h) for h in rows[hdr_at]]
    col = {h.upper(): i for i, h in enumerate(hdr) if h}
    kind = {i: classify(h) for i, h in enumerate(hdr) if h}
    income = {i: k for i, k in kind.items() if k in STREAMS}
    disc_cols = [i for i, k in kind.items() if k == 'discount']
    gst_cols = [i for i, k in kind.items() if k == 'gst']
    i_round, i_gross, i_type = col.get('ROUND OFF'), col['GROSS TOTAL'], col['VOUCHER TYPE']
    i_qty, i_val = col['QUANTITY'], col['VALUE']
    body, stated = [], None
    for r in rows[hdr_at + 1:]:
        if r[1] == 'Grand Total':
            stated = num(r[i_gross])
        else:
            body.append(r)

    mon = collections.defaultdict(lambda: collections.defaultdict(float))
    stream = collections.defaultdict(float)
    heads = collections.defaultdict(float)
    vt = collections.defaultdict(lambda: [0, 0.0, 0.0])
    day = collections.defaultdict(lambda: [0, 0, 0.0])
    items = collections.defaultdict(lambda: [0, 0.0, 0, collections.Counter()])
    sold = collections.defaultdict(lambda: collections.defaultdict(float))  # part no -> month -> qty
    cur, cur_fy, cur_m, flipped, odd = None, None, None, 0, 0
    for r in body:
        if r[0] is not None:
            d, t = r[0], r[i_type]
            m = d.strftime('%Y-%m')
            fy = fy_of(d)
            plus = sum(num(r[i]) for i in income)
            # Discount ledgers come as positive amounts in most exports and negative in some: use the size.
            disc = sum(abs(num(r[i])) for i in disc_cols)
            gst = sum(num(r[i]) for i in gst_cols)
            g, ro = num(r[i_gross]), num(r[i_round]) if i_round is not None else 0
            res = g - (plus - disc + gst)
            fits = lambda x: abs(x - ro) < 0.02 or abs(x + ro) < 0.02
            # The export drops minus signs: a bill whose discount exceeds it shows a positive total. Flip the sign
            # only when the negated total reconciles with the bill's own components; a bill that reconciles neither
            # way (a rare Tally export quirk) keeps the sign Tally gave it, which is what its Grand Total counts.
            if not fits(res) and fits(-g - (plus - disc + gst)):
                g = -g; flipped += 1
            elif not fits(res):
                odd += 1
            taxable = plus - disc
            svc = base_type(t) in SERVICE
            M = mon[m]
            M['invoices'] += 1; M['gross'] += g; M['taxable'] += taxable; M['gst'] += gst; M['discount'] += disc
            if svc:
                M['service_jobs'] += 1; M['service_gross'] += g
            for i, grp in income.items():
                if num(r[i]):
                    stream[(m, grp)] += num(r[i]); heads[(fy, head_name(hdr[i]), grp)] += num(r[i])
            stream[(m, 'Discounts')] -= disc
            for i in disc_cols:
                if num(r[i]):
                    heads[(fy, hdr[i], 'Discounts')] -= abs(num(r[i]))
            V = vt[(fy, t)]; V[0] += 1; V[1] += g; V[2] += taxable
            D = day[d.strftime('%Y-%m-%d')]; D[0] += 1; D[1] += svc; D[2] += g
            cur, cur_fy, cur_m = base_type(t), fy, m
        elif r[1]:
            name, pn = split_item(r[1])
            I = items[(cur_fy, pn, name)]
            I[0] += num(r[i_qty]); I[1] += num(r[i_val]); I[2] += 1; I[3][cur] += 1
            if not is_transfer(cur):
                sold[clean(r[1])][cur_m] += num(r[i_qty])
                if pn:
                    sold['[' + pn + ']'][cur_m] += num(r[i_qty])

    def chan(c):
        n = sum(c.values())
        if sum(c[s] for s in SERVICE) >= n / 2:
            return 'Service'
        if sum(c[s] for s in c if is_transfer(s)) >= n / 2:
            return 'Branch transfer'
        return 'Counter sale'

    months = sorted(mon)
    fym = lambda m: fy_of(datetime.date(int(m[:4]), int(m[5:7]), 1))
    top = []
    for fy in sorted({k[0] for k in items}):  # each year's top 150 items by value
        top += sorted([kv for kv in items.items() if kv[0][0] == fy], key=lambda kv: -kv[1][1])[:150]
    out = {
        'monthly': [[fym(m), m, int(mon[m]['invoices']), int(mon[m]['service_jobs'])] +
                    [round(mon[m][k], 2) for k in ['gross', 'taxable', 'gst', 'discount', 'service_gross']] for m in months],
        'streams': [[fym(m), m, s, round(v, 2)] for (m, s), v in sorted(stream.items())],
        'heads': sorted([[fy, h, g, round(v, 2)] for (fy, h, g), v in heads.items()], key=lambda x: (x[0], -x[3])),
        'voucher_types': sorted([[fy, k, v[0], round(v[1], 2), round(v[2], 2)] for (fy, k), v in vt.items()], key=lambda x: (x[0], -x[3])),
        'daily': [[fym(k[:7]), k, v[0], v[1], round(v[2], 2)] for k, v in sorted(day.items())],
        'items': [[fy, pn or nm, nm, q, round(v, 2), l, chan(c)] for (fy, pn, nm), (q, v, l, c) in top],
    }
    dates = sorted(day)
    gross = round(sum(v[1] for v in vt.values()), 2)
    if stated is None or abs(gross - stated) > 1:
        raise SystemExit(f'{os.path.basename(path)}: bills add up to {gross:,.2f} but the Grand Total says {stated}')
    info = {'register_from': dates[0], 'register_to': dates[-1], 'fys': sorted({r[0] for r in out['monthly']}),
            'bills': sum(v[0] for v in vt.values()), 'gross': gross, 'flipped': flipped, 'odd': odd,
            'item_value': round(sum(v[1] for v in items.values()), 2), 'distinct_items': len(items),
            'columns': {h: kind[i] for i, h in enumerate(hdr) if h and kind[i] != 'meta'}}
    return out, info, sold


def stock_summary(path):
    """A Tally Stock Summary: bold stock-group rows, each followed by its indented items, then Grand Total.
    Items are told from groups by Tally's own formatting, since not every item carries a part number."""
    import openpyxl
    ws = openpyxl.load_workbook(open(path, 'rb')).worksheets[0]
    rows = list(ws.iter_rows())
    at = next(i for i, r in enumerate(rows) if r[0].value == 'Particulars')
    period = clean(rows[at][1].value)
    as_at = datetime.datetime.strptime(period.split(' to ')[-1], '%d-%b-%y').date().isoformat()
    groups, its, stated, g, sg = [], [], None, None, None
    body = [r for r in rows[at + 3:] if r[0].value not in (None, '')]
    ind = lambda r: r[0].alignment.indent or 0
    for k, r in enumerate(body):
        c = r[0]
        q, v = num(r[1].value), num(r[3].value)
        if c.value == 'Grand Total':
            stated = num(r[3].value) if r[3].value is not None else num(r[2].value)
            break
        if c.font.b and not c.alignment.indent:
            g, sg = clean(c.value).title(), None
            groups.append([g, q, v])
        elif not c.font.i and k + 1 < len(body) and ind(body[k + 1]) > ind(r) and body[k + 1][0].value != 'Grand Total':
            sg = clean(c.value)   # a sub-group (Canon's "CM 18 %"): plain, not italic like items, its lines indented deeper
        else:
            its.append([clean(c.value), g, q, v, category(g, sg)])
    total = round(sum(x[3] for x in its), 2)
    if stated is None or abs(total - stated) > 1:
        raise SystemExit(f'{os.path.basename(path)}: items add up to {total:,.2f} but the Grand Total says {stated}')
    for gname, q, v in groups:  # every group row must equal its own items
        s = sum(x[3] for x in its if x[1] == gname)
        if abs(s - v) > 1:
            raise SystemExit(f'{os.path.basename(path)}: group {gname} says {v:,.2f}, its items add up to {s:,.2f}')
    return as_at, its, total


def stock_rows(as_at, its, sold):
    """Totals per stock group, and the 200 costliest lines held, each with how much of it sold in the
    register months of the same year up to as_at (counter, workshop and accessory sales; branch transfers
    excluded). A line held but not sold at all in those months counts as idle."""
    end, fy = as_at[:7], fy_of(datetime.date.fromisoformat(as_at))
    months = sorted({m for s in sold.values() for m in s
                     if m <= end and fy_of(datetime.date(int(m[:4]), int(m[5:7]), 1)) == fy})
    cats = collections.defaultdict(lambda: [0, 0.0, 0.0, 0, 0.0])
    detail = []
    for key, g, q, v, cat in its:
        name, pn = split_item(key)
        # Same name as in the register, else the same bracketed code (names are sometimes edited).
        hist = sold.get(key) or sold.get('[' + pn + ']', {}) if pn else sold.get(key, {})
        s = sum(hist.get(m, 0) for m in months)
        C = cats[cat or g or 'Ungrouped']; C[0] += 1; C[1] += q; C[2] += v
        if months and s <= 0 and v > 0:
            C[3] += 1; C[4] += v
        rate = s / len(months) if months else 0
        detail.append([as_at, pn or name, name, cat or g, q, round(v, 2), s if months else '', len(months) or '',
                       round(q / rate, 1) if rate > 0 else ''])
    detail.sort(key=lambda x: -x[5])
    cat_rows = [[as_at, c, n, q, round(v, 2), n2 if months else '', round(v2, 2) if months else '']
                for c, (n, q, v, n2, v2) in sorted(cats.items(), key=lambda kv: -kv[1][2])]
    return cat_rows, detail[:200], len(months)


def write_csv(path, cols, rows):
    with open(path, 'w', newline='') as f:
        w = csv.writer(f)
        w.writerow(cols)
        w.writerows(rows)


def read_csv(path):
    if not os.path.exists(path):
        return []
    with open(path, newline='') as f:
        return list(csv.DictReader(f))


def merge(out):
    companies = []
    for meta_path in sorted(glob.glob(os.path.join(out, 'companies', '*', 'company.json'))):
        with open(meta_path) as f:
            companies.append(json.load(f))
    for fid, cols in FILES.items():
        rows = []
        for c in companies:
            for r in read_csv(os.path.join(out, 'companies', c['code'], fid + '.csv')):
                rows.append([c['code']] + [r[k] for k in cols])
        write_csv(os.path.join(out, fid + '.csv'), ['company'] + cols, rows)
    write_csv(os.path.join(out, 'companies.csv'),
              ['company', 'name', 'register_from', 'register_to', 'bills', 'gross', 'fys'],
              [[c['code'], c['name'], c['register_from'], c['register_to'], c['bills'], c['gross'], ' '.join(c.get('fys', []))] for c in companies])
    received = {(r['company'], r['file']) for r in read_csv(os.path.join(out, 'received.csv'))}
    comp = []
    for c in companies:
        for report, period, stem in REQUIRED:
            key = f'{report} {period}'
            got = key in c.get('received', []) or (c['code'], key) in received
            comp.append([c['code'], key, report, period, f"{c['code']}_{stem}.xlsx", 'Received' if got else 'Pending'])
    write_csv(os.path.join(out, 'completeness.csv'), ['company', 'file', 'report', 'period', 'filename', 'status'], comp)
    return companies


def build(paths, code, name, out):
    """Summarise every file one company sent (any mix of sales registers and stock summaries)."""
    regs, stocks = [], []
    for p in paths:
        title, rows = sheet_rows(p)
        if any(r and r[0] == 'Date' and r[1] == 'Particulars' for r in rows[:40]):
            regs.append((p, rows))
        elif title == 'Stock Summary' or any(r and r[0] == 'Stock Summary' for r in rows[:15]):
            stocks.append((p, rows))
        else:
            raise SystemExit(f'{os.path.basename(p)}: neither a register (Date | Particulars header) nor a stock summary')
    if not regs:
        raise SystemExit('at least one sales register is needed')
    data = {fid: [] for fid in FILES}
    sold = collections.defaultdict(lambda: collections.defaultdict(float))
    infos, received, seen = [], [], set()
    for p, rows in regs:
        d, info, s = summarise(p, rows)
        clash = seen & set(info['fys'])
        if clash:
            raise SystemExit(f'{os.path.basename(p)}: FY {", ".join(clash)} is already covered by another register')
        seen |= set(info['fys'])
        for fid in d:
            data[fid] += d[fid]
        for k, v in s.items():
            for m, q in v.items():
                sold[k][m] += q
        infos.append(info)
        for fy in info['fys']:
            received.append('Sales register FY 2026-27 Apr-Sep' if fy == '2026-27' else f'Sales register FY {fy}')
        print(f"  {os.path.basename(p)}: {info['bills']:,} bills {info['register_from']} to {info['register_to']}, "
              f"gross {info['gross']:,.2f} = Grand Total ({info['flipped']} unsigned credit bills turned negative, {info.get('odd', 0)} kept as exported)")
        odd = {h: k for h, k in info['columns'].items() if k in ('Other sales', 'Branch transfers', 'AMC & warranty', 'Accessories')}
        print('    columns outside parts and labour:', odd)
    stock_meta = []
    for p, rows in stocks:
        as_at, its, total = stock_summary(p)
        cats, detail, months = stock_rows(as_at, its, sold)
        data['stock'] += cats
        data['stock_items'] += detail
        dt = datetime.date.fromisoformat(as_at)
        received.append(f"Stock summary As at {dt.day} {dt.strftime('%b %Y')}")
        stock_meta.append({'as_at': as_at, 'value': total, 'lines': len(its), 'sales_months': months})
        print(f"  {os.path.basename(p)}: stock at {as_at}, {len(its):,} lines, {total:,.2f} = Grand Total; "
              f"{months} months of sales to compare against")
    for fid in ('monthly', 'daily'):
        data[fid].sort(key=lambda r: r[1])
    d = os.path.join(out, 'companies', code)
    os.makedirs(d, exist_ok=True)
    for fid, cols in FILES.items():
        write_csv(os.path.join(d, fid + '.csv'), cols, data[fid])
    first = min(infos, key=lambda i: i['register_from'])
    last = max(infos, key=lambda i: i['register_to'])
    meta = {'code': code, 'name': name, 'register_from': first['register_from'], 'register_to': last['register_to'],
            'fys': sorted(seen), 'bills': sum(i['bills'] for i in infos), 'gross': round(sum(i['gross'] for i in infos), 2),
            'registers': [{k: i[k] for k in ('register_from', 'register_to', 'bills', 'gross', 'distinct_items')} for i in infos],
            'stock': sorted(stock_meta, key=lambda x: x['as_at']), 'received': sorted(set(received))}
    with open(os.path.join(d, 'company.json'), 'w') as f:
        json.dump(meta, f, indent=2)
    print(f"{code}: {meta['bills']:,} bills, gross {meta['gross']:,.2f}, FY {', '.join(meta['fys'])}; "
          f"{len(stock_meta)} stock summar{'y' if len(stock_meta) == 1 else 'ies'}")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument('files', nargs='*', help="all of one company's Tally exports: sales registers and stock summaries")
    ap.add_argument('--company', help='agreed company code, e.g. PC01 (letters, digits, - and _)')
    ap.add_argument('--name', help='company name as shown on the dashboard')
    ap.add_argument('--out', default='data')
    ap.add_argument('--profile', choices=['honda', 'canon'], default='honda', help='the business the registers come from')
    ap.add_argument('--merge-only', action='store_true', help='only rebuild the group files')
    a = ap.parse_args()
    use_profile(a.profile)
    if not a.merge_only:
        if not (a.files and a.company and a.name):
            ap.error('files, --company and --name are required unless --merge-only')
        if not re.fullmatch(r'[A-Za-z0-9_-]{1,32}', a.company):
            ap.error('--company may only use letters, digits, - and _')
        build(a.files, a.company, a.name, a.out)
    companies = merge(a.out)
    print(f"group files rebuilt in {a.out}/ for {len(companies)} compan{'y' if len(companies) == 1 else 'ies'}")


if __name__ == '__main__':
    main()
